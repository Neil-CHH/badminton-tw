# -*- coding: utf-8 -*-
"""解析官方「總成績紀錄」PDF 的成績總表 → standings(source=pdf)。

為什麼一定要用 PDF:推導名次推不出「該組取幾名」。實測 257164 新羽盃同一場裡
取 1 名／取 2 名／取 4 名(含並列第三)三種都有,籤表形狀卻一樣 —— 取幾名是
競賽規程的決定,比分資料裡沒有這個資訊。

mylivescore 產的 PDF 是可抽文字的向量表格(不是掃描檔),成績總表版面固定:

    ['項目', '第一名', '第二名', '第三名', '第三名', ...]   ← 名次表頭(並列就重複)
    [組別名,  單位1,    單位2,    單位3,    單位4  , ...]   ← 單位列
    ['',      選手1,    選手2,    選手3,    選手4  , ...]   ← 選手列(團體賽沒這列)

校驗原則:名次以 PDF 為準(官方文件),單位以比分資料為準(PDF 200dpi 小字常認錯)。
組別名完全相同就直接採信;只有「模糊配對來的組別」而且大半選手都對不上時,
才判定配錯而整組不收(寧缺勿錯)。比分沒收錄的組別/選手照 PDF 收,並在報告中標注。

用法:
    python -X utf8 scripts/parse_result_pdf.py --openid 257164            # 只看報告
    python -X utf8 scripts/parse_result_pdf.py --openid 257164 --apply    # 寫入
    python -X utf8 scripts/parse_result_pdf.py --all                      # 全庫掃描報告
    python -X utf8 scripts/parse_result_pdf.py --all --apply              # 全庫套用
"""
import argparse
import difflib
import json
import re
import sys
import urllib.parse
import urllib.request
from collections import defaultdict
from pathlib import Path

import fitz

ROOT = Path(__file__).resolve().parent.parent
TOURN_DIR = ROOT / "docs" / "data" / "tournaments"
sys.path.insert(0, str(ROOT / "scripts"))
from rebuild_index import split_members  # noqa: E402
from scrape import base_group  # noqa: E402

CACHE = ROOT / "inbox" / "resultpdf"          # inbox 已 gitignore
RESULT_KEYWORDS = ("總成績", "成績紀錄", "成績記錄", "成績總表", "成績冊", "名次表")

CN_NUM = {"一": 1, "二": 2, "三": 3, "四": 4, "五": 5,
          "六": 6, "七": 7, "八": 8, "九": 9, "十": 10}
RANK_WORD = {"冠軍": 1, "亞軍": 2, "季軍": 3, "殿軍": 4}
_RANK_RE = re.compile(r"^第([一二三四五六七八九十\d]+)名$")
_NORM_RE = re.compile(r"[\s()()\[\]【】‧・、,,\-－_/]")
PLACEHOLDER = ("單位", "選手一", "選手二", "隊名", "隊伍")
# 姓名切分:官方 PDF 也用全形斜線,rebuild_index.NAME_SPLIT 沒收這個字元
_EXTRA_SPLIT = re.compile(r"[／・]")


def split_names(cell):
    out = []
    for part in _EXTRA_SPLIT.split(cell or ""):
        out.extend(split_members(part))
    return [n for n in (clean_name(x) for x in out) if n]


def strip_unit_prefix(name, units):
    """姓名格裡混進單位名就剝掉。回傳 (修好的姓名, 是否有修)。

    682391 日月潭盃「公開國小高年級組男單」第二名,官方 PDF 的姓名格直接寫成
    「南投縣內湖國民小學龔品丞」(同一份的縣內組寫的是正確的「龔品丞」)——
    是主辦填表填錯,不是版面問題,剝掉前綴就對了。

    **不能改用字數上限或賽制字眼判斷姓名合不合理** —— 實測會誤殺真資料:原住民姓名
    「伊斯瑪哈撒嗯拉嘎夫」9 個字,而隊名本來就常含賽制字眼
    (「可以靠553拿名次嗎」「單打手的驕傲」「臺中羽球單打團」)。
    """
    for u in sorted((x for x in units if x and x != name), key=len, reverse=True):
        if name.startswith(u) and len(name) > len(u):
            return name[len(u):], True
        if name.endswith(u) and len(name) > len(u):
            return name[:-len(u)], True
    return name, False


def parse_rank(text):
    """'第三名' / '第 3 名' / '季軍' → 3;不是名次就回 None。"""
    t = re.sub(r"\s+", "", text or "")
    if t in RANK_WORD:
        return RANK_WORD[t]
    m = _RANK_RE.match(t)
    if not m:
        return None
    v = m.group(1)
    return int(v) if v.isdigit() else CN_NUM.get(v)


_ANNOT_RE = re.compile(r"[((][ㄅ-ㄩˊˇˋ˙\s]+[))]")   # 姓名旁的注音註記


def clean(s):
    """儲存格文字正規化。空白壓成一個半形空格而不是全刪 —— 全刪會把
    'Taguchi Fuminari' 併成 'TaguchiFuminari'(262259 實例)。"""
    return re.sub(r"\s+", " ", (s or "").replace("\n", " ")).strip()


def clean_name(s):
    """選手名:去掉注音註記(289558「林彥町（ㄊㄧㄥˇ）」),中文名再去掉內部空白。"""
    s = _ANNOT_RE.sub("", s or "").strip()
    if not re.search(r"[A-Za-z]", s):
        s = s.replace(" ", "")
    return s


def fetch_pdf(openid, url):
    CACHE.mkdir(parents=True, exist_ok=True)
    p = CACHE / f"{openid}.pdf"
    if p.exists() and p.stat().st_size > 1024:
        return p.read_bytes()
    req = urllib.request.Request(urllib.parse.quote(url, safe=":/%?=&"))
    req.add_header("User-Agent", "Mozilla/5.0 (badminton-db parse_result_pdf)")
    with urllib.request.urlopen(req, timeout=90) as res:
        data = res.read()
    p.write_bytes(data)
    return data


def result_pdf_url(t):
    for d in t.get("documents") or []:
        if any(k in (d.get("title") or "") for k in RESULT_KEYWORDS):
            return d.get("url"), d.get("title")
    return None, None


def event_arity(group):
    """組別名 → 一筆名次該有幾個選手名。團體 0、雙打 2、單打 1、判不出來 None。

    只在「完全沒有比分可以對照」時才用(排名賽/全大運的 API 回空)。組別名確實編碼了
    賽制,這跟拿賽制字眼去驗「人名合不合理」是兩回事 —— 後者會誤殺真隊名。
    """
    g = re.sub(r"\s+", "", group or "")
    if not g:
        return None
    if "團" in g or "隊" in g:
        return 0
    if "雙" in g:
        return 2
    if "單" in g:
        return 1
    return None


def cell_lines(raw):
    return [x.strip() for x in (raw or "").replace("\r", "").split("\n") if x.strip()]


def join_wrapped(lines):
    """把儲存格內被換行斷開的一段文字接回去。中文之間不補空格
    (「臺中市北屯區四維國民小」+「學」是單位名被排版折行,不是兩個欄位),
    英文之間才補,免得 'Taguchi' + 'Fuminari' 黏成一個字。"""
    out = ""
    for x in lines:
        if out and re.search(r"[A-Za-z0-9]$", out) and re.match(r"[A-Za-z0-9]", x):
            out += " "
        out += x
    return out


def parse_summary(doc, all_units=(), all_players=()):
    """回傳 [(組別, [(rank, 單位, [選手,...]), ...]), ...];找不到成績總表回 []。

    `all_units` / `all_players` 取自該賽事的比分資料,用來分辨「一格裡有換行」的兩種
    情況 —— 兩者長得一模一樣,只能靠實際資料判斷:
      A. 單位名太長被排版折行:「臺中市北屯區四維國民小」+「學」→ 接回去就是單位名
         (團體組很常見,341170/765740/535938 共 22 組都是這種)
      B. 單位與選手擠在同一格:「合庫北市大」+「王子維」→ 最後一行是選手
         (排名賽/全大運的版面)
    分不出來的組別回 (None, 組別名),交還人工,不猜。
    """
    out = []
    for page in doc:
        # 不能只認「成績總表」四個字:334198 全國國小盃的成績表沒寫這個標題,
        # 而 253124 南投市長盃整份都是籤表、根本沒有成績表。改認表格結構:
        # 第一欄是「項目」且同列有名次字樣,才是成績總表。
        text = page.get_text()
        if "項目" not in text:
            continue
        for tab in page.find_tables().tables:
            scan_table(tab.extract(), all_units, all_players, out)
    return finish_summary(out)


def scan_table(rows, all_units, all_players, out):
    """掃一張表格,把找到的成績總表區塊追加進 out。

    版面是「項目列 → 單位列 →(選手列)」三列一組,同一張表裡可以出現很多組
    (組別多到一頁排不下時,官方會再起一個項目列接著排)。
    """
    i = 0
    while i < len(rows):
        row = [clean(c) for c in rows[i]]
        if not row or row[0].replace(" ", "") != "項目":
            i += 1
            continue
        ranks = [(j, parse_rank(row[j])) for j in range(1, len(row))]
        ranks = [(j, r) for j, r in ranks if r]
        if not ranks or i + 1 >= len(rows):
            i += 1
            continue
        group = join_wrapped(cell_lines(rows[i + 1][0]))
        # 選手列:下一列第 0 欄是空的(組別欄用 rowspan 併格)才算
        name_row = None
        if i + 2 < len(rows):
            nxt = [clean(c) for c in rows[i + 2]]
            if not nxt[0] and any(nxt[1:]):
                name_row = nxt
        entries, unresolved = [], False
        for j, rank in ranks:
            unit, names = "", []
            if name_row and j < len(name_row) and name_row[j]:
                unit = join_wrapped(cell_lines(rows[i + 1][j])) \
                    if j < len(rows[i + 1]) else ""
                names = split_names(name_row[j])
            elif j < len(rows[i + 1]):
                lines = cell_lines(rows[i + 1][j])
                if len(lines) <= 1:
                    unit = join_wrapped(lines)
                else:
                    joined = join_wrapped(lines)
                    # 首選判準:最後一行是不是「這場比賽出現過的選手」。
                    # 是 → 版面 B(單位+選手同格);不是 → 單位名被折行。
                    # 不能反過來用「接起來是不是已知單位」判斷 —— 官方 PDF 寫
                    # 全稱「臺中市北屯區四維國民小學」,比分只寫「四維國小」,
                    # 對不上,團體組會全數落空(341170/765740/535938 共 22 組)。
                    if lines[-1] in all_players:
                        unit = join_wrapped(lines[:-1])
                        names = split_names(lines[-1])
                    elif joined in all_players:
                        names = split_names(joined)
                    elif all_players:
                        unit = joined      # 有名單可查卻查無 → 折行的單位名
                    else:
                        # 這場完全沒有比分可對照(排名賽/全大運 API 回空)。
                        # 退而用組別名判斷賽制:單打拿 1 個名字、雙打 2 個、
                        # 團體 0 個。不硬猜 —— 判不出賽制就交還人工。
                        k = event_arity(group)
                        if k is None or len(lines) <= k:
                            # 組別名看不出賽制(「女校長組」「教育行政人員組」)。
                            # 先當成折行的單位名記著,等整份表看完再決定 ——
                            # 若整份表沒有任何一組出現選手名,那就是純團體賽事,
                            # 這些多行儲存格必然是隊名折行(282076 教育盃)。
                            unresolved = True
                            unit = joined
                        elif k == 0:
                            unit = joined
                        else:
                            unit = join_wrapped(lines[:-k])
                            for ln in lines[-k:]:
                                names.extend(split_names(ln))
            if unit in PLACEHOLDER or any(n in PLACEHOLDER for n in names):
                continue                      # 官方沒填,只留樣板字
            if not unit and not names:
                continue                      # 空格/畫斜線 = 從缺
            entries.append((rank, unit, names))
        if group and entries:
            out.append((None if unresolved else group, entries, group))
        i += 3 if name_row else 2


def finish_summary(out):
    # 純團體賽事的補救:整份成績表都沒有選手名 → 判不出賽制的那幾組就是隊名折行
    if not any(ns for g, es, _n in out if g is not None for _r, _u, ns in es):
        out = [(g if g is not None else name, es, name) for g, es, name in out]
    return [(g, es) if g is not None else (None, name) for g, es, name in out]


def parse_summary_xlsx(path, all_units=(), all_players=()):
    """xlsx 版的成績總表。版面與 PDF 完全一樣(主辦是同一套表格另存成 xlsx),
    所以只把儲存格讀成 rows 再交給同一支 scan_table。"""
    import openpyxl
    out = []
    wb = openpyxl.load_workbook(path, data_only=True)
    for ws in wb.worksheets:
        rows = [["" if c is None else str(c) for c in r]
                for r in ws.iter_rows(values_only=True)]
        scan_table(rows, all_units, all_players, out)
    return finish_summary(out)


def resolve_group(pdf_group, known, roster, pdf_names):
    """PDF 組別名 → (賽事組別名, 對照結果)。

    純比字串會出事:277843 彰化縣長盃的比分只有國小 12 組,PDF 卻有 25 組(含國中、
    高中)。「國中女生組單打」跟「國小中年級女生組單打」字面相似度夠高,硬配就會把
    國中組的獎掛到國小組。所以模糊比對必須**再用選手名驗證**:候選組的比分名單要真的
    出現 PDF 上的人,才算對上;都對不上就當成「比分沒收錄的新組別」照 PDF 收錄。
    """
    if pdf_group in known:
        return pdf_group, "exact"
    target = _NORM_RE.sub("", pdf_group)
    for k in known:
        if _NORM_RE.sub("", k) == target:
            return k, "exact"
    cands = sorted(
        ((difflib.SequenceMatcher(None, target, _NORM_RE.sub("", k)).ratio(), k)
         for k in known), reverse=True)
    for score, k in cands:
        if score < 0.75:
            break
        names = roster.get(k) or {}
        if not pdf_names or not names:
            continue
        if pdf_names & set(names):          # 名字對得上才承認這個配對
            return k, "fuzzy"
    return pdf_group, "new"                 # 比分沒有這組 → 照 PDF 原名收錄


def match_roster(t):
    """賽事比分資料 → {組別: {選手: {單位,...}}},用來校驗 PDF 判讀。"""
    roster = defaultdict(lambda: defaultdict(set))
    for m in t.get("matches") or []:
        g = base_group(m.get("groupName", ""))
        for side, team in (("A", m.get("teamA")), ("B", m.get("teamB"))):
            for si in m.get("scoreinfo") or []:
                for n in split_members(si.get("member" + side) or ""):
                    roster[g][n].add(team or "")
    return roster


def manual_openids():
    """逐列人工判讀過的賽事(import_pdf_standings + import_observation_standings)。

    那些 PDF 常用本程式不支援的版面,而且多半沒有比分可以交叉驗證(全大運/排名賽的
    API 回空),硬解析只會把已經正確的資料洗爛 —— **整場跳過,不是只跳過同名組別**
    (實測漏掉「同一場多出來的組別」就足以寫進亂資料)。
    """
    out = set()
    for mod in ("import_pdf_standings", "import_observation_standings"):
        try:
            out |= set(__import__(mod).DATA)
        except Exception:             # noqa: BLE001 匯入失敗就當作沒有人工資料
            pass
    # 全運會/全中運同理:掛在 mylivescore 的「成績紀錄」PDF 其實是完成後的籤表,
    # 名次已由主辦競賽資訊系統的官方頒獎名單補上(scrape_sportgov),不必再解 PDF。
    try:
        import scrape_sportgov as _sg
        out |= set(_sg.SITES) | set(_sg.QUALIFIERS)
    except Exception:                 # noqa: BLE001
        pass
    return out


# 官方成績總表不在 mylivescore 的賽事:mylivescore 掛的「總成績紀錄」其實是完成後的
# 籤表(有「取 N 名」但沒有成績總表),真正的成績總表由主辦另外公告,只能人工取得。
# 專案不留底官方 PDF,但這兩份是「唯一取得管道」——外站沒有穩定連結,所以收進 Ref/。
LOCAL_SUMMARY = {
    "537403": "Ref/114年臺中久樘盃國小羽球錦標賽_總成績表.pdf",
    "751804": "Ref/2026金門縣浯洲盃全國羽球邀請賽_成績總表.xlsx",
}

MANUAL = manual_openids()


def process(openid, apply=False, local=None):
    if openid in MANUAL:
        return {"openid": openid, "status": "名次已由其他來源匯入,跳過"}
    path = TOURN_DIR / f"{openid}.json"
    t = json.loads(path.read_text(encoding="utf-8"))
    local = local or LOCAL_SUMMARY.get(openid)
    if not local:
        url, _title = result_pdf_url(t)
        if not url:
            return {"openid": openid, "status": "無總成績PDF"}

    roster = match_roster(t)
    all_units = {u for u in
                 ((m.get(k) or "").strip() for m in t.get("matches") or []
                  for k in ("teamA", "teamB")) if len(u) >= 2}
    all_players = {n for g in roster.values() for n in g}
    try:
        if local:
            src = local if Path(local).is_absolute() else ROOT / local
            if not src.exists():
                return {"openid": openid, "status": f"找不到本地成績總表:{local}"}
            if src.suffix.lower() in (".xlsx", ".xlsm"):
                parsed = parse_summary_xlsx(src, all_units, all_players)
            else:
                parsed = parse_summary(fitz.open(src), all_units, all_players)
        else:
            parsed = parse_summary(
                fitz.open(stream=fetch_pdf(openid, url), filetype="pdf"),
                all_units, all_players)
    except Exception as e:            # noqa: BLE001 下載/解析失敗只記錄,不中斷整批
        return {"openid": openid, "status": f"PDF失敗:{type(e).__name__}"}
    # 版面不支援是「以組別為單位」的問題:同一份 PDF 常常絕大多數組別都是支援的版面,
    # 只有少數幾組(或別的附表)把單位與選手擠在同一格。整場放棄會白白丟掉大量正確資料
    # —— 341170 臺中市長盃 2,692 場比分就是這樣整場被誤判掉的。
    unsupported = [g2 for g1, g2 in parsed if g1 is None]
    parsed = [(g, e) for g, e in parsed if g is not None]
    if not parsed:
        return {"openid": openid,
                "status": ("版面不支援(單位與選手同格),需人工" if unsupported
                           else "找不到成績總表")}


    known = [k for k in
             sorted({base_group(m.get("groupName", "")) for m in t.get("matches") or []}
                    | {g.get("name") for g in t.get("groups") or [] if g.get("name")})
             if k]
    accepted, rejected, warned, fixes = [], [], [], 0
    new_groups = 0
    for pdf_group, entries in parsed:
        pdf_names = {n for _r, _u, ns in entries for n in ns}
        g, how = resolve_group(pdf_group, known, roster, pdf_names)
        in_group = roster.get(g, {})
        if how == "new":
            new_groups += 1
        # 姓名格裡混進單位名 → 剝掉前綴修好(官方填表失誤,見 strip_unit_prefix)
        cleaned_entries, stripped = [], []
        for rank, unit, names in entries:
            fixed_names = []
            for n in names:
                fixed, did = strip_unit_prefix(n, all_units)
                if did:
                    stripped.append(f"{n}→{fixed}")
                fixed_names.append(fixed)
            cleaned_entries.append((rank, unit, fixed_names))
        entries = cleaned_entries
        pdf_names = {n for _r, _u, ns in entries for n in ns}
        if stripped:
            warned.append((pdf_group, "姓名格混進單位名,已剝除:" + "; ".join(stripped[:3])))
        rows, unmatched = [], []
        for rank, unit, names in entries:
            for n in names:
                if in_group and n not in in_group:
                    unmatched.append(f"第{rank}名「{n}」")
            fixed = unit
            if names and in_group:      # 單位以比分為準(PDF 小字常認錯)
                units = set()
                for n in names:
                    units |= {u for u in in_group.get(n, set()) if u}
                if len(units) == 1:
                    fixed = next(iter(units))
            if fixed != unit:
                fixes += 1
            rows.append({"group": g, "rank": rank, "unit": fixed,
                         "members": names, "source": "pdf"})
        # 名字大半對不上,只有在「組別是模糊配對來的」時才代表配錯 → 整組不收。
        # 組名完全相同就是同一組,名字對不上只表示該組比分收得不全(官方 PDF 才是全的),
        # 這時照收 PDF,否則會白白丟掉正確的名次。
        named = [n for _r, _u, ns in entries for n in ns]
        if how == "fuzzy" and named and in_group and len(unmatched) / len(named) > 0.5:
            rejected.append((pdf_group, f"→{g}:{len(unmatched)}/{len(named)} 名對不上,"
                                        f"疑似組別配錯"))
            continue
        if unmatched:
            warned.append((g, "; ".join(unmatched[:4])))
        accepted.extend(rows)

    res = {"openid": openid, "name": t.get("name"), "status": "OK",
           "groups": len({r["group"] for r in accepted}), "rows": len(accepted),
           "unitFixes": fixes, "newGroups": new_groups,
           "unsupported": unsupported,
           "rejected": rejected, "warned": warned,
           "before": sorted({s.get("source") for s in t.get("standings") or []})}

    if apply and accepted:
        pdf_groups = {r["group"] for r in accepted}
        kept = [s for s in t.get("standings") or [] if s.get("group") not in pdf_groups]
        t["standings"] = sorted(kept + accepted, key=lambda s: (s["group"], s["rank"]))
        existing = {g.get("name") for g in t.get("groups") or []}
        for g in sorted(pdf_groups - existing):
            t.setdefault("groups", []).append(
                {"id": "", "name": g, "tags": [], "drawUrl": None})
        path.write_text(json.dumps(t, ensure_ascii=False, separators=(",", ":")),
                        encoding="utf-8")
        res["applied"] = True
    return res


def targets_all(force=False):
    """有官方總成績 PDF 的賽事。預設跳過「名次已全是 pdf」的(省下重跑),
    但解析規則改過之後要用 --force 全部重跑,否則舊結果不會更新。"""
    out = []
    for p in sorted(TOURN_DIR.glob("*.json")):
        t = json.loads(p.read_text(encoding="utf-8"))
        srcs = {s.get("source") for s in t.get("standings") or []}
        if srcs == {"pdf"} and not force:
            continue
        if result_pdf_url(t)[0] or p.stem in LOCAL_SUMMARY:
            out.append(p.stem)
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--openid")
    ap.add_argument("--file", help="改讀本地成績總表(pdf/xlsx),不下載官方連結")
    ap.add_argument("--all", action="store_true")
    ap.add_argument("--apply", action="store_true")
    ap.add_argument("--force", action="store_true",
                    help="連已經全是 pdf 名次的賽事也重跑(改過解析規則後要用)")
    ap.add_argument("--json")
    args = ap.parse_args()

    targets = ([args.openid] if args.openid
               else (targets_all(args.force) if args.all else None))
    if not targets:
        ap.error("需要 --openid 或 --all")
    if args.file and len(targets) != 1:
        ap.error("--file 只能配 --openid 用")

    results, ok, fail, rows, rej, warn, newg = [], 0, 0, 0, 0, 0, 0
    for i, oid in enumerate(targets, 1):
        r = process(oid, apply=args.apply, local=args.file)
        results.append(r)
        if r["status"] == "OK":
            ok += 1
            rows += r["rows"]
            rej += len(r["rejected"])
            warn += len(r["warned"])
            newg += r["newGroups"]
            flag = "!" if r["rejected"] else ("+" if r["newGroups"] else " ")
            print(f"[{i}/{len(targets)}]{flag} {oid:<12} {r['groups']:>3} 組 "
                  f"{r['rows']:>4} 筆  {(r.get('name') or '')[:26]}")
            for g, why in r["rejected"][:3]:
                print(f"        [不收] {g}:{why}")
            if r.get("unsupported"):
                print(f"        [版面不支援] {len(r['unsupported'])} 組需人工:"
                      f"{r['unsupported'][:3]}")
            for g, why in r["warned"][:2]:
                print(f"        [注意] {g}:{why}(比分未收錄該選手,照 PDF 收)")
        else:
            fail += 1
            print(f"[{i}/{len(targets)}]  {oid:<12} {r['status']}")
    print(f"\n完成:成功 {ok} 場、共 {rows} 筆名次;比分沒有的新組別 {newg} 組;"
          f"零星選手對不上 {warn} 組;不收 {rej} 組;無法處理 {fail} 場")
    if args.apply:
        print("已寫入賽事檔 → 接著跑 python scripts/rebuild_index.py")
    if args.json:
        Path(args.json).write_text(
            json.dumps(results, ensure_ascii=False, indent=2, default=list),
            encoding="utf-8")


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    main()
